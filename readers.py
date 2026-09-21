"""Reading layer: turn procurement files into text and tables.

This module owns everything about getting content out of files, kept separate
from the graph logic in extract_graph.py. It handles nested archives, several
file formats, an OCR fallback for scanned PDFs, and table extraction.

read_file() returns text and tables together. PDFs are checked page by page,
so image-only pages inside a digital PDF also receive OCR. Short readable text
is kept. read_text() and extract_tables() remain for callers needing one side.

The base readers (pdfplumber, python-docx, openpyxl) are imported at the top.
OCR and table models load only when needed. A table-model failure is reported
but does not prevent page OCR or discard previously recovered content.
"""

from __future__ import annotations

import io
import re
import time
import zipfile
import shutil
import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory
from html.parser import HTMLParser

# Base readers, used on every run, are imported here as usual. The optional OCR
# (rapidocr, onnxruntime) and Table Transformer (transformers, torch) paths are
# imported inside the functions that use them, so the CPV pipeline starts quickly
# without them installed.
import docx
import openpyxl
import pdfplumber


# CONFIGURATION

CONTAINER_EXTS = (".zip", ".edoc")       # archive types to open instead of read as text
MAX_CONTAINER_DEPTH = 5                  # stop runaway recursion on nested archives
MIN_USEFUL_CHARS = 50                    # OCR sparse PDF pages that also contain images

IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp")
IMAGE_MAX_EDGE = 2400                    # near the long edge of A4 at 200 DPI
OCR_STRIP_CJK = True                      # Latvian corpus only; disable for multilingual input
_CJK_IDEOGRAPHS = re.compile("[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\U00020000-\U000323af]+")
DOC_CONVERSION_TIMEOUT = 60

OCR_DPI = 200                            # resolution for rasterising scanned PDF pages

ENABLE_TABLE_TRANSFORMER = True
TATR_DETECTION_MODEL = "microsoft/table-transformer-detection"
TATR_STRUCTURE_MODEL = "microsoft/table-structure-recognition-v1.1-all"
TATR_THRESHOLD = 0.7                     # min confidence for a detected table/row/column
RAPIDOCR_MODEL_TYPE = "small"            # PP-OCRv6 tier for the page OCR ('small' or 'medium')


# ARCHIVES

def iter_container_files(data, container_name, depth=0):
    """Yield (filename, bytes) for every leaf file inside a ZIP or .edoc archive.

    Names retain the archive chain. Signature metadata inside .edoc containers
    is skipped. Invalid archive bytes are yielded unchanged as one leaf file.
    """
    if depth >= MAX_CONTAINER_DEPTH:
        return
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for info in archive.infolist():
                if info.is_dir():
                    continue
                name = info.filename
                relative = name.replace("\\", "/").lstrip("./").lower()
                if container_name.lower().endswith(".edoc") and (
                        relative == "mimetype" or relative.startswith("meta-inf/")):
                    continue
                source = f"{container_name} > {name}"
                payload = archive.read(info)
                if name.lower().endswith(CONTAINER_EXTS):
                    yield from iter_container_files(payload, source, depth + 1)
                else:
                    yield source, payload
    except zipfile.BadZipFile:
        yield container_name, data


# TEXT READERS
#
# read_plain is the digital-only reader used by the comparison tools.
# read_text also handles scanned PDF pages through the combined reading path.

def _doc_to_docx(data):
    """Convert legacy Word bytes locally, keeping the original file unchanged."""
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        mac_path = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        if mac_path.is_file():
            executable = str(mac_path)
    if not executable:
        raise RuntimeError("Install LibreOffice to read .doc files (soffice not found)")
    with TemporaryDirectory(prefix="procurement_doc_") as folder:
        folder = Path(folder)
        source = folder / "input.doc"
        source.write_bytes(data)
        # A separate profile avoids interference with an open LibreOffice window.
        result = subprocess.run(
            [executable, f"-env:UserInstallation={(folder / 'profile').as_uri()}",
             "--headless", "--convert-to", "docx", "--outdir", str(folder), str(source)],
            capture_output=True, text=True, timeout=DOC_CONVERSION_TIMEOUT,
        )
        converted = folder / "input.docx"
        if result.returncode != 0 or not converted.is_file():
            detail = (result.stderr or result.stdout).strip()[:300]
            raise RuntimeError(f"DOC conversion failed: {detail or 'no DOCX produced'}")
        return converted.read_bytes()


def _image_pages(data):
    """Decode image frames, applying camera orientation before OCR."""
    from PIL import Image, ImageOps, ImageSequence
    with Image.open(io.BytesIO(data)) as source:
        images = []
        for frame in ImageSequence.Iterator(source):
            image = ImageOps.exif_transpose(frame).convert("RGB")
            image.thumbnail((IMAGE_MAX_EDGE, IMAGE_MAX_EDGE), Image.Resampling.LANCZOS)
            images.append(image)
        return images


def read_plain(data, name):
    """Extract text from PDF, DOCX, XLSX, or TXT bytes by file extension."""
    lower = name.lower()
    try:
        if lower.endswith(".pdf"):
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        if lower.endswith(".doc"):
            data = _doc_to_docx(data)
            lower = ".docx"
        if lower.endswith(".docx"):
            return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
        if lower.endswith((".xlsx", ".xlsm")):
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append("\t".join(cells))
            return "\n".join(rows)
        if lower.endswith((".txt", ".csv")):
            return data.decode("utf-8", errors="ignore")
    except Exception as error:
        print(f"    could not read {name}: {error}")
    return ""


def _pdf_page_images(data, page_numbers=None):
    """Render selected zero-based PDF pages, or all pages when omitted."""
    import fitz  # PyMuPDF
    from PIL import Image
    images = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        indexes = range(len(doc)) if page_numbers is None else page_numbers
        for index in indexes:
            page = doc[index]
            pix = page.get_pixmap(dpi=OCR_DPI)
            images.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
    return images


_rapidocr_reader = None


def _get_rapidocr_reader():
    """Build and cache the RapidOCR engine (PP-OCRv6, covers Latvian) once."""
    global _rapidocr_reader
    if _rapidocr_reader is None:
        from rapidocr import RapidOCR, ModelType, OCRVersion
        _rapidocr_reader = RapidOCR(params={"Rec.ocr_version": OCRVersion("PP-OCRv6"),
                                            "Rec.model_type": ModelType(RAPIDOCR_MODEL_TYPE)})
    return _rapidocr_reader


def _clean_ocr_text(text):
    """Suppress CJK ideograph noise for this Latvian corpus, preserving token gaps."""
    if OCR_STRIP_CJK:
        text = _CJK_IDEOGRAPHS.sub(" ", text)
    return text.strip()


def _ocr_texts(reader, image):
    """Run RapidOCR on an image; return the list of recognised text pieces."""
    import numpy as np
    result = reader(np.array(image))
    texts = getattr(result, "txts", None)               # rapidocr v3 result object
    if texts:
        texts = list(texts)
    elif isinstance(result, tuple) and result[0]:      # older (result, elapse) form
        texts = [row[1] for row in result[0]]
    else:
        return []
    return [cleaned for text in texts if (cleaned := _clean_ocr_text(text))]


def _ocr_cell(reader, crop):
    """Read one table-cell image with RapidOCR; return its text (pieces joined)."""
    return " ".join(_ocr_texts(reader, crop))


def read_ocr(data, name, images=None):
    """Recover text from PDF pages or image files with RapidOCR.

    Args:
      data: Raw PDF or image bytes.
      name: File name; must end in .pdf or a supported image extension.
      images: Optional pre-rendered pages to reuse instead of rasterising again.

    Returns:
      The recognised text, or "" with an install hint when RapidOCR is missing.
    """
    if not name.lower().endswith((".pdf",) + IMAGE_EXTS):
        return ""
    try:
        reader = _get_rapidocr_reader()            # RapidOCR (PP-OCRv6)
        if images is None:
            images = (_pdf_page_images(data) if name.lower().endswith(".pdf")
                      else _image_pages(data))
        lines = []
        for image in images:
            lines.extend(_ocr_texts(reader, image))
        return "\n".join(lines)
    except ImportError as error:
        print(f"    skipping OCR for {name}: run 'pip install rapidocr onnxruntime pymupdf' to enable it ({error})")
        return ""
    except Exception as error:
        print(f"    OCR failed for {name}: {error}")
        return ""


TEXT_READERS = [read_plain, read_ocr]


def read_text(data, name):
    """Keep readable text, using page-level OCR for PDFs without table models."""
    if name.lower().endswith((".pdf", ".doc") + IMAGE_EXTS):
        return read_file(data, name, include_tables=False)[0]
    for reader in TEXT_READERS:
        text = reader(data, name)
        if text and text.strip():
            return text
    return ""


# TABLE EXTRACTION
#
# Each returned table is {"source": filename, "rows": [[cell, ...], ...]}.
# Presence-checking is built in: a file with no tables yields an empty list.

def _tables_from_docx(data):
    """Read Word tables as lists of rows. Word exposes them directly."""
    document = docx.Document(io.BytesIO(data))
    tables = []
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            tables.append(rows)
    return tables


def _tables_from_xlsx(data):
    """Treat each worksheet as one table of rows."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cell.strip() for cell in cells):
                rows.append(cells)
        if rows:
            tables.append(rows)
    return tables


def _tables_from_pdf(data, name):
    """Read tables from digital and scanned pages without full-page OCR."""
    return _read_pdf(data, name, {}, include_text=False)[1]


def extract_tables(data, name):
    """Return a list of tables found in one file; empty when there are none."""
    lower = name.lower()
    try:
        if lower.endswith(".doc"):
            data = _doc_to_docx(data)
            lower = ".docx"
        if lower.endswith(".docx"):
            rows_list = _tables_from_docx(data)
        elif lower.endswith((".xlsx", ".xlsm")):
            rows_list = _tables_from_xlsx(data)
        elif lower.endswith(".pdf"):
            rows_list = _tables_from_pdf(data, name)
        elif lower.endswith(IMAGE_EXTS) and ENABLE_TABLE_TRANSFORMER:
            rows_list = _tables_from_images(_image_pages(data))[0]
        else:
            rows_list = []
    except Exception as error:
        print(f"    table extraction failed for {name}: {error}")
        return []
    return [{"source": name, "rows": rows} for rows in rows_list]


# TABLE TRANSFORMER
#
# Two DETR-based models. The pipeline uses only the detection model, to locate table
# regions that RapidTable then reconstructs; the structure model and _cells_to_rows
# stay for compare_tables.py.

_tatr_processor = None
_tatr_detection = None
_tatr_structure = None


def _load_tatr(include_structure=True):
    """Load detection, plus structure when requested by comparison tools."""
    global _tatr_processor, _tatr_detection, _tatr_structure
    if _tatr_detection is None:
        from transformers import AutoImageProcessor, TableTransformerForObjectDetection
        _tatr_processor = AutoImageProcessor.from_pretrained(TATR_DETECTION_MODEL)
        _tatr_detection = TableTransformerForObjectDetection.from_pretrained(TATR_DETECTION_MODEL)
    if include_structure and _tatr_structure is None:
        from transformers import TableTransformerForObjectDetection
        _tatr_structure = TableTransformerForObjectDetection.from_pretrained(TATR_STRUCTURE_MODEL)
    return _tatr_processor, _tatr_detection, _tatr_structure


def _detect_objects(processor, model, image):
    """Run a Table Transformer model on one image; return [{label, score, box}].

    Uses the processor's post_process_object_detection, which handles the
    box maths and drops the "no object" class for us (much less code than doing
    it by hand).
    """
    import torch
    inputs = processor(images=image, return_tensors="pt")
    with torch.no_grad():
        outputs = model(**inputs)
    target_sizes = torch.tensor([image.size[::-1]])
    result = processor.post_process_object_detection(
        outputs, threshold=TATR_THRESHOLD, target_sizes=target_sizes
    )[0]
    objects = []
    for score, label, box in zip(result["scores"], result["labels"], result["boxes"]):
        objects.append({
            "label": model.config.id2label[int(label)],
            "score": float(score),
            "box": [int(v) for v in box.tolist()],
        })
    return objects


def _cells_to_rows(structure, image):
    """Turn detected rows and columns into text rows by OCR-ing each cell with RapidOCR."""
    rows = sorted((o for o in structure if o["label"] == "table row"),
                  key=lambda o: o["box"][1])           # top to bottom
    cols = sorted((o for o in structure if o["label"] == "table column"),
                  key=lambda o: o["box"][0])           # left to right
    reader = _get_rapidocr_reader()
    table = []
    for row in rows:
        line = []
        for col in cols:
            # A cell is the intersection of one row band and one column band.
            cell_box = (col["box"][0], row["box"][1], col["box"][2], row["box"][3])
            line.append(_ocr_cell(reader, image.crop(cell_box)))
        table.append(line)
    return table


_rapidtable_engine = None


def _get_rapidtable_engine():
    """Build and cache the RapidTable engine (SLANet structure + RapidOCR cells) once."""
    global _rapidtable_engine
    if _rapidtable_engine is None:
        from rapid_table import RapidTable
        _rapidtable_engine = RapidTable()
    return _rapidtable_engine


class _TableParser(HTMLParser):
    """Collect HTML cells with decoded text and their row/column spans."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows = []
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
            self.rows.append(self.row)
        elif tag in ("td", "th") and self.row is not None:
            attrs = dict(attrs)
            spans = []
            for key in ("rowspan", "colspan"):
                try:
                    value = int(attrs.get(key, "1"))
                except (ValueError, TypeError):
                    value = 1
                spans.append(value if value >= 1 or (key == "rowspan" and value == 0) else 1)
            self.cell = {"text": [], "rowspan": spans[0], "colspan": spans[1]}
            self.row.append(self.cell)
        elif tag in ("br", "p", "div") and self.cell is not None:
            self.cell["text"].append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self.cell = None
        elif tag == "tr":
            self.cell = None
            self.row = None
        elif tag in ("p", "div") and self.cell is not None:
            self.cell["text"].append(" ")

    def handle_data(self, data):
        if self.cell is not None:
            self.cell["text"].append(data)


def _html_to_rows(html_text):
    """Expand merged cells to a rectangular grid using empty covered cells.

    Text stays in the top-left cell of a span, so it is not counted repeatedly
    downstream. Empty placeholders keep subsequent values in their columns.
    """
    parser = _TableParser()
    parser.feed(html_text)
    parser.close()
    grid = {}
    width = 0
    for row_index, cells in enumerate(parser.rows):
        column = 0
        for cell in cells:
            colspan = cell["colspan"]
            while any((row_index, column + offset) in grid for offset in range(colspan)):
                column += 1
            # rowspan=0 means through the remaining rows of this table.
            rowspan = cell["rowspan"] or len(parser.rows) - row_index
            for row in range(row_index, min(row_index + rowspan, len(parser.rows))):
                for col in range(column, column + colspan):
                    grid[row, col] = ""
            grid[row_index, column] = " ".join("".join(cell["text"]).split())
            column += colspan
            width = max(width, column)
    return [[grid.get((row, col), "") for col in range(width)]
            for row in range(len(parser.rows))] if width else []


def _mask_regions(images, boxes_per_page):
    """Copy each page with its table boxes painted white, so the full-page OCR skips them."""
    from PIL import ImageDraw
    out = []
    for image, boxes in zip(images, boxes_per_page):
        if boxes:
            image = image.copy()
            draw = ImageDraw.Draw(image)
            for box in boxes:
                draw.rectangle(box, fill="white")
        out.append(image)
    return out


def _tables_from_images(images):
    """Detect table regions with the Table Transformer and reconstruct them with RapidTable.

    Returns:
      (tables, table_boxes_per_page). A box is reported only when RapidTable returns a
      table for it, so the caller can mask captured regions from the full-page OCR while
      a failed region is left for the OCR.
    """
    import numpy as np
    processor, detection, _ = _load_tatr(include_structure=False)
    engine = _get_rapidtable_engine()
    tables, boxes_per_page = [], []
    for image in images:
        page_boxes = []
        for obj in _detect_objects(processor, detection, image):
            if obj["label"] not in ("table", "table rotated"):
                continue
            box = tuple(obj["box"])
            try:
                htmls = engine(np.asarray(image.crop(box))).pred_htmls
                rows = _html_to_rows(htmls[0]) if htmls else []
                rows = [[_clean_ocr_text(cell) for cell in row] for row in rows]
            except Exception as error:
                print(f"    table recognition failed on a region: {error}")
                rows = []
            if rows:
                tables.append(rows)
                page_boxes.append(box)                # mask only regions we captured
        boxes_per_page.append(page_boxes)
    return tables, boxes_per_page


def _tables_from_scanned_pdf(data):
    """Render a scanned PDF's pages and recognise tables on them (renders here)."""
    tables, _ = _tables_from_images(_pdf_page_images(data))
    return tables


# COMBINED READER

def _pdf_text_and_tables(data, include_tables=True, warnings=None):
    """Return per-page text, digital tables and zero-based pages needing OCR.

    Image-only pages need OCR. A page with sparse text and an image also needs
    it, since its text may be only a header above a scan. Short digital text is
    kept even when no OCR is needed. Table errors do not discard page text.
    """
    text_parts, rows_list = [], []
    scan_pages = []
    warnings = warnings if warnings is not None else []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for index, page in enumerate(pdf.pages):
            try:
                page_text = page.extract_text() or ""
            except Exception as error:
                warnings.append(f"page {index + 1} text: {error}")
                page_text = ""
            text_parts.append(page_text)
            if not page_text.strip() or (len(page_text.strip()) < MIN_USEFUL_CHARS and page.images):
                scan_pages.append(index)
            if include_tables:
                try:
                    for table in page.extract_tables():
                        rows = [["" if c is None else str(c) for c in row] for row in table]
                        if rows:
                            rows_list.append(rows)
                except Exception as error:
                    warnings.append(f"page {index + 1} tables: {error}")
    return text_parts, rows_list, scan_pages


def _read_pdf(data, name, info, include_text=True, include_tables=True):
    """Read each PDF page, reusing its rendered image for tables and OCR."""
    warnings = info.setdefault("warnings", [])
    page_texts, tables, scan_pages = _pdf_text_and_tables(data, include_tables, warnings)
    info["ocr_seconds"] = 0.0
    info["ocr_text"] = ""
    info["ocr_page_numbers"] = []
    ocr_parts = []
    if scan_pages and (include_text or (include_tables and ENABLE_TABLE_TRANSFORMER)):
        try:
            images = _pdf_page_images(data, scan_pages)
        except Exception as error:
            warnings.append(f"page rendering: {error}")
            images = []
        for index, image in zip(scan_pages, images):
            ocr_images = [image]
            if include_tables and ENABLE_TABLE_TRANSFORMER:
                try:
                    captured, boxes = _tables_from_images([image])
                    tables.extend(captured)
                    ocr_images = _mask_regions([image], boxes)
                except Exception as error:
                    warnings.append(f"page {index + 1} scanned tables: {error}")
                    # Keep the unmasked image available to OCR.
            if include_text:
                started = time.perf_counter()
                ocr = read_ocr(data, name, ocr_images)
                info["ocr_seconds"] += time.perf_counter() - started
                if ocr.strip():
                    ocr_parts.append(ocr)
                    info["ocr_page_numbers"].append(index + 1)
                    page_texts[index] = "\n".join(part for part in (page_texts[index], ocr) if part)
                else:
                    warnings.append(f"page {index + 1}: OCR returned no text")
    info["ocr_text"] = "\n".join(ocr_parts)
    for warning in warnings:
        print(f"    {name}: {warning}")
    return "\n".join(page_texts) if include_text else "", tables


def _docx_text_and_tables(data):
    """Read a DOCX's paragraph text and its tables."""
    document = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)
    rows_list = []
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            rows_list.append(rows)
    return text, rows_list


def _xlsx_text_and_tables(data):
    """Read an XLSX's cell text and per-sheet tables in one pass (read-only sheets iterate once)."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    text_lines, rows_list = [], []
    for sheet in workbook.worksheets:
        sheet_rows = []
        for row in sheet.iter_rows(values_only=True):
            text_cells = [str(c) for c in row if c is not None]
            if text_cells:
                text_lines.append("\t".join(text_cells))
            table_cells = ["" if c is None else str(c) for c in row]
            if any(cell.strip() for cell in table_cells):
                sheet_rows.append(table_cells)
        if sheet_rows:
            rows_list.append(sheet_rows)
    return "\n".join(text_lines), rows_list


def read_file(data, name, info=None, *, include_tables=True):
    """Read a file into text and tables, checking PDF pages individually.

    Args:
      data: The file's raw bytes.
      name: File name; its extension selects the reader.
      info: Optional dict, filled in for the caller to report on the read, with keys
        reader ("pdf"/"doc"/"docx"/"xlsx"/"text"/"ocr", or None if unread), reason,
        parse_seconds, ocr_seconds, ocr_text, ocr_page_numbers and warnings.
      include_tables: False for the text-only entry point.

    Returns:
      (text, tables). Short text is preserved. PDF text combines the digital
      content with recovered OCR in page order. info["ocr_text"] keeps only
      OCR content for the separate artifact. Table spans use empty placeholders.
    """
    if info is None:
        info = {}                 # local scratch when the caller passed none
    info["reader"] = None
    info["reason"] = ""
    info["parse_seconds"] = 0.0
    info["ocr_seconds"] = 0.0
    info["ocr_text"] = ""
    info["ocr_page_numbers"] = []
    info["warnings"] = []

    lower = name.lower()
    started = time.perf_counter()
    try:
        if lower.endswith(".pdf"):
            plain, rows_list = _read_pdf(data, name, info, include_tables=include_tables)
            fmt = "pdf"
        elif lower.endswith(".doc"):
            plain, rows_list = _docx_text_and_tables(_doc_to_docx(data))
            fmt = "doc"
        elif lower.endswith(IMAGE_EXTS):
            images = _image_pages(data)
            rows_list = []
            ocr_images = images
            if include_tables and ENABLE_TABLE_TRANSFORMER:
                try:
                    rows_list, boxes = _tables_from_images(images)
                    ocr_images = _mask_regions(images, boxes)
                except Exception as error:
                    info["warnings"].append(f"image tables: {error}")
                    print(f"    image table extraction failed for {name}: {error}")
            ocr_started = time.perf_counter()
            plain = read_ocr(data, name, ocr_images)
            info["ocr_seconds"] = time.perf_counter() - ocr_started
            info["ocr_text"] = plain
            fmt = "ocr"
        elif lower.endswith(".docx"):
            plain, rows_list = _docx_text_and_tables(data)
            fmt = "docx"
        elif lower.endswith((".xlsx", ".xlsm")):
            plain, rows_list = _xlsx_text_and_tables(data)
            fmt = "xlsx"
        elif lower.endswith((".txt", ".csv")):
            plain, rows_list = data.decode("utf-8", errors="ignore"), []
            fmt = "text"
        else:
            ext = name.lower().rsplit(".", 1)[-1] if "." in name else "?"
            info["reason"] = f"unsupported type (.{ext})"
            return "", []
    except Exception as error:
        info["reason"] = f"{type(error).__name__}: {error}"
        print(f"    could not read {name}: {error}")
        return "", []
    finally:
        info["parse_seconds"] = round(max(0.0, time.perf_counter() - started - info["ocr_seconds"]), 2)
        info["ocr_seconds"] = round(info["ocr_seconds"], 2)

    if plain.strip():
        info["reader"] = "ocr" if info["ocr_text"] else fmt
    else:
        info["reason"] = "no text recovered"
    tables = [{"source": name, "rows": rows} for rows in rows_list] if include_tables else []
    return plain, tables
